"""
DOCX Parser using python-docx.
Extracts clean Markdown and structured tables from Microsoft Word (.docx) documents.
"""

import os
from typing import Tuple, List
import docx
from pipelines.schema import TableData
from pipelines.extractors.table_utils import format_matrix_to_markdown_table, extract_markdown_tables


class DocxParser:
    """Parses DOCX documents into clean Markdown and structured TableData objects."""

    def parse(self, file_path: str) -> Tuple[str, List[TableData]]:
        """
        Parses a .docx file into clean Markdown and structured TableData objects.

        Args:
            file_path: Path to DOCX file.

        Returns:
            Tuple of (clean_markdown_string, list_of_TableData)
        """
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"DOCX file not found: {file_path}")

        doc = docx.Document(file_path)
        md_chunks: List[str] = []
        tables_data: List[TableData] = []
        table_counter = 1

        # Iterate through paragraphs and tables in document order
        # doc.element.body contains paragraph and table xml elements in order
        for child in doc.element.body:
            if child.tag.endswith("p"):
                # It's a paragraph
                p_elem = child
                # Find matching paragraph in doc.paragraphs
                p_obj = None
                for p in doc.paragraphs:
                    if p._p == p_elem:
                        p_obj = p
                        break

                if p_obj:
                    text = p_obj.text.strip()
                    if not text:
                        continue
                    style_name = (p_obj.style.name or "").lower()

                    if "heading 1" in style_name:
                        md_chunks.append(f"# {text}\n")
                    elif "heading 2" in style_name:
                        md_chunks.append(f"## {text}\n")
                    elif "heading 3" in style_name:
                        md_chunks.append(f"### {text}\n")
                    elif "heading 4" in style_name:
                        md_chunks.append(f"#### {text}\n")
                    elif "list" in style_name or "bullet" in style_name:
                        md_chunks.append(f"- {text}")
                    else:
                        md_chunks.append(f"{text}\n")

            elif child.tag.endswith("tbl"):
                # It's a table
                tbl_elem = child
                tbl_obj = None
                for t in doc.tables:
                    if t._tbl == tbl_elem:
                        tbl_obj = t
                        break

                if tbl_obj:
                    extracted_rows: List[List[str]] = []
                    for row in tbl_obj.rows:
                        row_text = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                        extracted_rows.append(row_text)

                    if extracted_rows:
                        headers = extracted_rows[0]
                        data_rows = extracted_rows[1:] if len(extracted_rows) > 1 else []
                        table_md = format_matrix_to_markdown_table(headers, data_rows)

                        md_chunks.append("\n" + table_md + "\n")
                        tables_data.append(TableData(
                            table_index=table_counter,
                            headers=headers,
                            rows=data_rows,
                            markdown_representation=table_md,
                            row_count=len(data_rows),
                            column_count=len(headers)
                        ))
                        table_counter += 1

        full_markdown = "\n".join(md_chunks).strip()
        # Also check if any tables were missed
        if not tables_data:
            tables_data = extract_markdown_tables(full_markdown)

        return full_markdown, tables_data
