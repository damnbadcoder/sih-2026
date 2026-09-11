"""
DOCX Parser using python-docx and Tesseract OCR.
Extracts clean Markdown, structured tables, and embedded images/screenshots
from Microsoft Word (.docx) documents.
"""

import os
from typing import Tuple, List
import docx
from pipelines.text_pipeline.schema import TableData
from pipelines.text_pipeline.extractors.table_utils import format_matrix_to_markdown_table, extract_markdown_tables
from pipelines.text_pipeline.extractors.ocr_utils import extract_text_from_image_bytes, is_ocr_available


class DocxParser:
    """Parses DOCX documents into clean Markdown, structured TableData objects, and OCR image context."""

    def parse(self, file_path: str) -> Tuple[str, List[TableData]]:
        """
        Parses a .docx file into clean Markdown and structured TableData objects.
        Extracts embedded figures and screenshots via OCR.

        Args:
            file_path: Path to DOCX file.

        Returns:
            Tuple of (clean_markdown_string, list_of_TableData)
        """
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"DOCX file not found: {file_path}")

        doc = docx.Document(file_path)
        md_chunks: List[str] = []
        tables_data: List[TableData] = []
        table_counter = 1

        # Iterate through paragraphs and tables in document order
        for child in doc.element.body:
            if child.tag.endswith("p"):
                p_elem = child
                p_obj = None
                for p in doc.paragraphs:
                    if p._p == p_elem:
                        p_obj = p
                        break

                if p_obj:
                    text = p_obj.text.strip()
                    if not text:
                        continue
                    style_name = (p_obj.style.name or "").lower()

                    if "heading 1" in style_name:
                        md_chunks.append(f"# {text}\n")
                    elif "heading 2" in style_name:
                        md_chunks.append(f"## {text}\n")
                    elif "heading 3" in style_name:
                        md_chunks.append(f"### {text}\n")
                    elif "heading 4" in style_name:
                        md_chunks.append(f"#### {text}\n")
                    elif "list" in style_name or "bullet" in style_name:
                        md_chunks.append(f"- {text}")
                    else:
                        md_chunks.append(f"{text}\n")

            elif child.tag.endswith("tbl"):
                tbl_elem = child
                tbl_obj = None
                for t in doc.tables:
                    if t._tbl == tbl_elem:
                        tbl_obj = t
                        break

                if tbl_obj:
                    extracted_rows: List[List[str]] = []
                    for row in tbl_obj.rows:
                        row_text = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                        extracted_rows.append(row_text)

                    if extracted_rows:
                        headers = extracted_rows[0]
                        data_rows = extracted_rows[1:] if len(extracted_rows) > 1 else []
                        table_md = format_matrix_to_markdown_table(headers, data_rows)

                        md_chunks.append("\n" + table_md + "\n")
                        tables_data.append(TableData(
                            table_index=table_counter,
                            headers=headers,
                            rows=data_rows,
                            markdown_representation=table_md,
                            row_count=len(data_rows),
                            column_count=len(headers)
                        ))
                        table_counter += 1

        # Extract embedded images and screenshots via OCR
        image_ocr_blocks = self._extract_embedded_images_ocr(doc)
        if image_ocr_blocks:
            md_chunks.append("\n## 🖼️ Extracted Embedded Image & Screenshot Telemetry (OCR)\n")
            md_chunks.extend(image_ocr_blocks)

        full_markdown = "\n".join(md_chunks).strip()
        # Fallback table check
        if not tables_data:
            tables_data = extract_markdown_tables(full_markdown)

        return full_markdown, tables_data

    def _extract_embedded_images_ocr(self, doc: docx.Document) -> List[str]:
        """Scans document relationships for embedded graphics and extracts text using OCR."""
        if not is_ocr_available():
            return []

        ocr_blocks: List[str] = []
        image_idx = 1

        try:
            for rel_id, part in doc.part.related_parts.items():
                if hasattr(part, "content_type") and part.content_type.startswith("image/"):
                    try:
                        img_bytes = part.blob
                        ocr_text = extract_text_from_image_bytes(img_bytes)
                        if ocr_text:
                            ocr_blocks.append(
                                f"### 🖼️ Embedded Graphic / Screenshot {image_idx} (OCR)\n\n"
                                f"```text\n{ocr_text}\n```\n"
                            )
                            image_idx += 1
                    except Exception:
                        continue
        except Exception:
            pass

        return ocr_blocks
