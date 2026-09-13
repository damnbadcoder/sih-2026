import io
import uuid
from unittest.mock import AsyncMock
import pytest
from pptx import Presentation
from app.models.input_file import InputFile
from app.processing.inspection.audio import AudioInspector
from app.processing.inspection.pptx import PPTXInspector


def create_mock_file(filename: str, size: int) -> InputFile:
    return InputFile(
        id=uuid.uuid4(),
        job_id=uuid.uuid4(),
        original_filename=filename,
        stored_filename=filename,
        content_type="application/octet-stream",
        file_size=size,
        storage_path=f"test/{filename}",
    )


def create_mock_storage(data: bytes):
    storage = AsyncMock()

    async def _read_chunks(_path):
        yield data

    storage.read_chunks = _read_chunks
    return storage


@pytest.mark.asyncio
async def test_pptx_text_and_notes_extraction():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Forensic Analysis Title"

    notes_slide = slide.notes_slide
    notes_slide.notes_text_frame.text = "Confidential presenter notes"

    buffer = io.BytesIO()
    prs.save(buffer)
    content = buffer.getvalue()

    mock_file = create_mock_file("sample.pptx", len(content))
    mock_storage = create_mock_storage(content)

    inspector = PPTXInspector()
    result = await inspector.inspect(mock_file, mock_storage)

    assert result.extracted_text is not None
    assert "Forensic Analysis Title" in result.extracted_text
    assert "Confidential presenter notes" in result.extracted_text
    assert result.metadata["slide_count"] == 1


@pytest.mark.asyncio
async def test_audio_empty_bytes_graceful_handling():
    content = b""
    mock_file = create_mock_file("sample.mp3", len(content))
    mock_storage = create_mock_storage(content)

    inspector = AudioInspector()
    result = await inspector.inspect(mock_file, mock_storage)

    assert result.extracted_text is not None
    assert "Audio Asset (sample.mp3)" in result.extracted_text
    assert result.metadata["duration_seconds"] == 0.0