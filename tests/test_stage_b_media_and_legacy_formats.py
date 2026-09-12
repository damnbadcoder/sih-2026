import io
import json
import struct
import uuid
import wave
from collections.abc import AsyncIterator
from pathlib import Path

import pytest
import xlwt
from httpx import ASGITransport, AsyncClient
from pptx import Presentation
from sqlalchemy import delete

from app.core.formats import (
    MEDIA_CATEGORY_AUDIO,
    MEDIA_CATEGORY_DOCUMENT,
    MEDIA_CATEGORY_IMAGE,
    MEDIA_CATEGORY_PRESENTATION,
    MEDIA_CATEGORY_TEXT,
    MEDIA_CATEGORY_VIDEO,
    classify_format,
)
from app.core.uploads import ALLOWED_UPLOAD_TYPES
from app.db.session import async_session_factory
from app.main import app
from app.models.input_file import InputFile
from app.models.job import Job
from app.models.user import User
from app.processing.inspection import (
    AudioInspector,
    EVTXInspector,
    ImageInspector,
    InspectionError,
    PPTXInspector,
    SVGInspector,
    VideoInspector,
    XLSInspector,
    get_inspector,
)
from app.processing.normalization import iter_normalized_records
from app.processing.service import ProcessingError, process_job, reserve_job_for_processing
from app.storage import get_storage
from app.storage.local import LocalStorage

_EVTX_FIXTURE = Path(__file__).parent / "fixtures" / "evtx" / "issue_38.evtx"
_EVTX_MALFORMED_FIXTURE = Path(__file__).parent / "fixtures" / "evtx" / "dns_log_malformed.evtx"

CREATED_EMAILS: list[str] = []


def unique_email(prefix: str) -> str:
    return f"{prefix}-{uuid.uuid4().hex[:8]}@example.com"


async def cleanup_created_users() -> None:
    if not CREATED_EMAILS:
        return
    async with async_session_factory() as db:
        await db.execute(delete(User).where(User.email.in_(CREATED_EMAILS)))
        await db.commit()
    CREATED_EMAILS.clear()


@pytest.fixture
def storage(tmp_path) -> LocalStorage:
    store = LocalStorage(tmp_path / "uploads")
    app.dependency_overrides[get_storage] = lambda: store
    yield store
    app.dependency_overrides.pop(get_storage, None)


@pytest.fixture
async def client(storage):
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as ac:
        yield ac
    await cleanup_created_users()


# --- minimal, spec-valid format builders --------------------------------------


def build_bmp(width: int = 320, height: int = 240, bit_depth: int = 24) -> bytes:
    pixel_offset = 14 + 40
    row_size = (width * bit_depth // 8 + 3) & ~3
    file_size = pixel_offset + row_size * height
    file_header = b"BM" + struct.pack("<IHHI", file_size, 0, 0, pixel_offset)
    info_header = struct.pack(
        "<IiiHHIIiiII", 40, width, height, 1, bit_depth, 0, 0, 0, 0, 0, 0
    )
    return file_header + info_header + b"\x00" * (row_size * height)


def build_webp(width: int = 320, height: int = 200) -> bytes:
    packed = (width - 1) | ((height - 1) << 14)
    payload = b"\x2f" + struct.pack("<I", packed)  # VP8L signature + lossless dims
    chunk = b"VP8L" + struct.pack("<I", len(payload)) + payload + b"\x00"
    return b"RIFF" + struct.pack("<I", 4 + len(chunk)) + b"WEBP" + chunk


def _tiff_entry(tag: int, field_type: int, value: int) -> bytes:
    return struct.pack("<HHII", tag, field_type, 1, value)


def build_tiff(width: int = 320, height: int = 240, bit_depth: int = 8) -> bytes:
    entries = (
        _tiff_entry(256, 4, width)  # ImageWidth (LONG)
        + _tiff_entry(257, 4, height)  # ImageLength (LONG)
        + _tiff_entry(258, 3, bit_depth)  # BitsPerSample (SHORT)
        + _tiff_entry(259, 3, 1)  # Compression = none
        + _tiff_entry(262, 3, 2)  # Photometric = RGB
    )
    ifd = struct.pack("<H", 5) + entries + struct.pack("<I", 0)
    return b"II" + struct.pack("<H", 42) + struct.pack("<I", 8) + ifd


def build_wav(sample_rate: int = 8000, seconds: int = 1, channels: int = 1) -> bytes:
    buf = io.BytesIO()
    with wave.open(buf, "wb") as writer:
        writer.setnchannels(channels)
        writer.setsampwidth(2)
        writer.setframerate(sample_rate)
        writer.writeframes(b"\x00\x00" * sample_rate * seconds)
    return buf.getvalue()


def build_flac(sample_rate: int = 44100, total_samples: int = 441000) -> bytes:
    value = (sample_rate << 44) | (1 << 41) | (15 << 36) | total_samples
    streaminfo = (
        struct.pack(">HH", 4096, 4096)
        + b"\x00\x00\x00\x00\x00\x00"  # min/max frame size unknown
        + value.to_bytes(8, "big")
        + b"\x00" * 16  # MD5 signature (not decoded)
    )
    block = b"\x80" + b"\x00\x00\x22" + streaminfo  # last | type 0 (STREAMINFO), len 34
    return b"fLaC" + block


def build_ogg(rate: int = 44100, total_samples: int = 441000) -> bytes:
    ident = (
        struct.pack("<I", 0)  # vorbis version
        + bytes([2])  # channels
        + struct.pack("<I", rate)  # sample rate
        + struct.pack("<I", 0)  # bitrate maximum
        + struct.pack("<I", 160000)  # bitrate nominal
        + struct.pack("<I", 0)  # bitrate minimum
        + bytes([0xB0])  # blocksizes
        + bytes([1])  # framing
    )
    first_payload = b"\x01vorbis" + ident
    first_page = (
        b"OggS\x00\x00"
        + struct.pack("<Q", 0)
        + struct.pack("<I", 1)
        + struct.pack("<I", 0)
        + struct.pack("<I", 0)
        + b"\x02\x07\x1e"
        + first_payload
    )
    final_page = (
        b"OggS\x00\x04"
        + struct.pack("<Q", total_samples)
        + struct.pack("<I", 1)
        + struct.pack("<I", 1)
        + struct.pack("<I", 0)
        + b"\x00"
    )
    return first_page + b"\x00" * 64 + final_page


def _ebml_size(n: int) -> bytes:
    if n < 0x80:
        return bytes([0x80 | n])
    if n < 0x4000:
        return bytes([0x40 | (n >> 8), n & 0xFF])
    if n < 0x200000:
        return bytes([0x20 | (n >> 16), (n >> 8) & 0xFF, n & 0xFF])
    return bytes([0x10 | (n >> 24), (n >> 16) & 0xFF, (n >> 8) & 0xFF, n & 0xFF])


def _ebml_element(eid: bytes, payload: bytes) -> bytes:
    return eid + _ebml_size(len(payload)) + payload


def build_mkv(doctype: bytes = b"matroska", duration: float = 5500.0) -> bytes:
    ebml = _ebml_element(
        b"\x1a\x45\xdf\xa3",
        _ebml_element(b"\x42\x86", b"\x01")
        + _ebml_element(b"\x42\x82", doctype)
        + _ebml_element(b"\x42\x87", b"\x02"),
    )
    info = _ebml_element(
        b"\x15\x49\xa9\x66",
        _ebml_element(b"\x2a\xd7\xb1", (1_000_000).to_bytes(4, "big"))
        + _ebml_element(b"\x44\x89", struct.pack(">d", duration)),
    )
    video = _ebml_element(
        b"\xe0", _ebml_element(b"\xb0", b"\x02\x80") + _ebml_element(b"\xba", b"\x01\xe0")
    )
    audio = _ebml_element(
        b"\xe1",
        _ebml_element(b"\xb5", struct.pack(">d", 44100.0)) + _ebml_element(b"\x9f", b"\x02"),
    )
    video_track = (
        _ebml_element(b"\xae", _ebml_element(b"\x83", b"\x01")
        + _ebml_element(b"\x86", b"V_VP9") + video)
    )
    audio_track = (
        _ebml_element(b"\xae", _ebml_element(b"\x83", b"\x02")
        + _ebml_element(b"\x86", b"A_OPUS") + audio)
    )
    tracks = _ebml_element(b"\x16\x54\xae\x6b", video_track + audio_track)
    segment = _ebml_element(b"\x18\x53\x80\x67", info + tracks)
    return ebml + segment


def _riff_chunk(cc: bytes, payload: bytes) -> bytes:
    pad = b"\x00" if len(payload) % 2 else b""
    return cc + struct.pack("<I", len(payload)) + payload + pad


def build_avi(width: int = 320, height: int = 240, frames: int = 1000) -> bytes:
    avih_payload = struct.pack(
        "<IIIIIIIIIIII", 40000, 0, 0, 0, frames, 0, 1, 0, width, height, 0, 0
    )
    strh_payload = (
        b"vids"
        + b"DIB "
        + struct.pack("<IHHIIIIIIII", 0, 0, 0, 0, 1, 25, 0, frames, 0, 0, 0)
        + struct.pack("<iiii", 0, 0, 0, 0)
    )
    strf_payload = struct.pack("<IiiHHIIiiII", 40, width, height, 1, 24, 0, 0, 0, 0, 0, 0)
    avih = _riff_chunk(b"avih", avih_payload)
    strh = _riff_chunk(b"strh", strh_payload)
    strf = _riff_chunk(b"strf", strf_payload)
    strl = b"LIST" + struct.pack("<I", 4 + len(strh) + len(strf)) + b"strl" + strh + strf
    hdrl = b"LIST" + struct.pack("<I", 4 + len(avih) + len(strl)) + b"hdrl" + avih + strl
    return b"RIFF" + struct.pack("<I", 4 + len(hdrl)) + b"AVI " + hdrl


def _box(tag: bytes, payload: bytes) -> bytes:
    return struct.pack(">I", 8 + len(payload)) + tag + payload


def build_mov_or_m4a(handler: bytes = b"vide") -> bytes:
    mvhd = _box(
        b"mvhd",
        b"\x00\x00\x00\x00" + struct.pack(">I", 0) + struct.pack(">I", 0)
        + struct.pack(">I", 1000) + struct.pack(">I", 3000),
    )
    hdlr = _box(b"hdlr", b"\x00\x00\x00\x00\x00\x00\x00\x00" + handler + b"\x00" * 12)
    mdia = _box(b"mdia", hdlr)
    trak = _box(b"trak", mvhd + mdia)
    return _box(b"ftyp", b"isom\x00\x00\x00\x00qt  ") + _box(b"moov", trak)


def build_pptx(lines: list[str]) -> bytes:
    prs = Presentation()
    blank = prs.slide_layouts[6]
    for line in lines:
        slide = prs.slides.add_slide(blank)
        textbox = slide.shapes.add_textbox(0, 0, 100, 100)
        textbox.text = line
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def build_xls() -> bytes:
    workbook = xlwt.Workbook()
    sheet = workbook.add_sheet("Alerts")
    sheet.write(0, 0, "host")
    sheet.write(0, 1, "count")
    sheet.write(1, 0, "srv-01")
    sheet.write(1, 1, 42)
    buf = io.BytesIO()
    workbook.save(buf)
    return buf.getvalue()


def build_svg(body: str, root: str = "svg") -> bytes:
    return (
        f'<?xml version="1.0" encoding="UTF-8"?>'
        f'<{root} xmlns="http://www.w3.org/2000/svg" width="640" height="480" '
        f'viewBox="0 0 640 480">{body}</{root}>'
    ).encode()


def read_evtx(path: Path) -> bytes:
    return path.read_bytes()


# --- helpers ------------------------------------------------------------------


async def _bytes_chunks(data: bytes) -> AsyncIterator[bytes]:
    yield data


def make_input_file(filename: str, content_type: str, data: bytes) -> InputFile:
    return InputFile(
        id=uuid.uuid4(),
        job_id=uuid.uuid4(),
        original_filename=filename,
        stored_filename=f"stored_{filename}",
        content_type=content_type,
        file_size=len(data),
        storage_path=f"{uuid.uuid4()}/job/{filename}",
    )


async def write_via_storage(storage: LocalStorage, input_file: InputFile, data: bytes) -> None:
    await storage.save(_bytes_chunks(data), input_file.storage_path, max_size=len(data) + 1)


async def _inspect(storage: LocalStorage, filename: str, data: bytes):
    suffix = Path(filename).suffix
    inspector = get_inspector(classify_format(filename, ALLOWED_UPLOAD_TYPES[suffix]))
    assert inspector is not None, filename
    input_file = make_input_file(filename, ALLOWED_UPLOAD_TYPES[suffix], data)
    await write_via_storage(storage, input_file, data)
    return await inspector.inspect(input_file, storage)


# --- format-aware selection ----------------------------------------------------


STAGE_B_FORMATS = {
    ".pptx": PPTXInspector,
    ".xls": XLSInspector,
    ".webp": ImageInspector,
    ".svg": SVGInspector,
    ".tiff": ImageInspector,
    ".tif": ImageInspector,
    ".bmp": ImageInspector,
    ".wav": AudioInspector,
    ".m4a": AudioInspector,
    ".ogg": AudioInspector,
    ".flac": AudioInspector,
    ".mkv": VideoInspector,
    ".mov": VideoInspector,
    ".avi": VideoInspector,
    ".webm": VideoInspector,
    ".evtx": EVTXInspector,
}


@pytest.mark.parametrize("extension", sorted(STAGE_B_FORMATS))
def test_stage_b_extension_resolves_registered_inspector(extension: str):
    inspector = get_inspector(
        classify_format(f"sample{extension}", ALLOWED_UPLOAD_TYPES[extension])
    )
    assert isinstance(inspector, STAGE_B_FORMATS[extension])


@pytest.mark.parametrize("extension", sorted(STAGE_B_FORMATS))
def test_stage_b_formats_carry_expected_media_category(extension: str):
    expected = {
        ".pptx": MEDIA_CATEGORY_PRESENTATION,
        ".xls": MEDIA_CATEGORY_DOCUMENT,
        ".webp": MEDIA_CATEGORY_IMAGE,
        ".svg": MEDIA_CATEGORY_IMAGE,
        ".tiff": MEDIA_CATEGORY_IMAGE,
        ".tif": MEDIA_CATEGORY_IMAGE,
        ".bmp": MEDIA_CATEGORY_IMAGE,
        ".wav": MEDIA_CATEGORY_AUDIO,
        ".m4a": MEDIA_CATEGORY_AUDIO,
        ".ogg": MEDIA_CATEGORY_AUDIO,
        ".flac": MEDIA_CATEGORY_AUDIO,
        ".mkv": MEDIA_CATEGORY_VIDEO,
        ".mov": MEDIA_CATEGORY_VIDEO,
        ".avi": MEDIA_CATEGORY_VIDEO,
        ".webm": MEDIA_CATEGORY_VIDEO,
        ".evtx": MEDIA_CATEGORY_TEXT,
    }
    classification = classify_format(f"sample{extension}", ALLOWED_UPLOAD_TYPES[extension])
    assert classification.media_category == expected[extension]


# --- extended images -----------------------------------------------------------


@pytest.mark.parametrize(
    ("filename", "builder", "expect"),
    [
        (
            "img.bmp",
            build_bmp,
            {"format": "bmp", "width": 320, "height": 240, "bit_depth": 24},
        ),
        (
            "img.webp",
            build_webp,
            {"format": "webp", "width": 320, "height": 200, "variant": "lossless"},
        ),
        (
            "img.tiff",
            build_tiff,
            {"format": "tiff", "width": 320, "height": 240, "bit_depth": 8},
        ),
    ],
)
async def test_extended_image_inspection_extracts_dimensions(
    storage: LocalStorage, filename, builder, expect
):
    result = await _inspect(storage, filename, builder())
    assert result.supported_for_inspection is True
    assert result.media_category == MEDIA_CATEGORY_IMAGE
    assert result.extracted_text is None
    for key, value in expect.items():
        assert result.metadata[key] == value


@pytest.mark.parametrize(
    ("filename", "bad_bytes"),
    [
        ("junk.bmp", b"not a bitmap at all"),
        ("junk.webp", b"not a webp container"),
        ("junk.tiff", b"II\x00\x00 not a tiff"),
        ("junk.svg", b"<div>this is not svg</div>"),
    ],
)
async def test_extended_image_inspection_rejects_garbage(
    storage: LocalStorage, filename, bad_bytes
):
    with pytest.raises(InspectionError):
        await _inspect(storage, filename, bad_bytes)


# --- SVG: data-only extraction ------------------------------------------------


async def test_svg_inspection_extracts_text_and_geometry(storage: LocalStorage):
    svg = build_svg('<title>Diagram</title><text>host 10.0.0.1</text><tspan>alert=true</tspan>')
    result = await _inspect(storage, "plot.svg", svg)

    assert result.supported_for_inspection is True
    assert result.media_category == MEDIA_CATEGORY_IMAGE
    assert result.metadata["width"] == "640"
    assert result.metadata["height"] == "480"
    assert result.metadata["view_box"] == "0 0 640 480"
    assert result.metadata["script_count"] == 0
    assert result.metadata["external_reference_count"] == 0
    assert result.extracted_text is not None
    assert "Diagram" in result.extracted_text
    assert "host 10.0.0.1" in result.extracted_text
    assert "alert=true" in result.extracted_text


async def test_svg_inspection_reports_scripts_but_never_executes(storage: LocalStorage):
    svg = build_svg(
        '<script>alert("pwned")</script><text>safe text</text>'
        '<script type="text/javascript">window.location = "https://evil.example"</script>'
    )
    result = await _inspect(storage, "scripted.svg", svg)

    assert result.metadata["script_count"] == 2
    assert result.extracted_text is not None
    lowered = result.extracted_text.lower()
    assert "alert(" not in lowered
    assert "safe text" in result.extracted_text


async def test_svg_inspection_reports_external_references_never_fetches(storage: LocalStorage):
    svg = build_svg(
        '<image href="http://evil.example/steal.png" width="10" height="10"/>'
        '<use href="#local-symbol"/><text>ok</text>'
    )
    result = await _inspect(storage, "linked.svg", svg)

    assert result.metadata["external_reference_count"] == 1
    assert result.extracted_text == "ok"


async def test_svg_inspection_rejects_non_svg_root(storage: LocalStorage):
    svg = build_svg("<p>html-ish</p>", root="html")
    with pytest.raises(InspectionError):
        await _inspect(storage, "fake.svg", svg)


async def test_svg_inspection_is_indicative_data_not_ocr(storage: LocalStorage):
    """SVG text is real element text, distinct from pixel/OCR extraction."""
    result = await _inspect(storage, "label.svg", build_svg("<text>alert-A</text>"))
    record = next(iter_normalized_records([result]))
    assert record["normalized_text"] is not None
    assert record["reason"] is None
    assert "alert-A" in record["normalized_text"]


# --- extended audio ------------------------------------------------------------


@pytest.mark.parametrize(
    ("filename", "builder", "expect_fmt"),
    [
        ("snd.wav", lambda: build_wav(8000, 1, 1), "wav"),
        ("snd.flac", lambda: build_flac(), "flac"),
        ("snd.ogg", lambda: build_ogg(), "ogg"),
    ],
)
async def test_extended_audio_inspection_extracts_header_metadata(
    storage: LocalStorage, filename, builder, expect_fmt
):
    result = await _inspect(storage, filename, builder())
    assert result.supported_for_inspection is True
    assert result.media_category == MEDIA_CATEGORY_AUDIO
    assert result.extracted_text is None
    assert result.metadata["format"] == expect_fmt
    assert result.metadata["sample_rate_hz"] == 44100 or result.metadata["sample_rate_hz"] == 8000
    assert result.metadata["duration_estimate_seconds"] is not None


async def test_audio_m4a_extracts_duration_and_audio_flow(storage: LocalStorage):
    result = await _inspect(storage, "clip.m4a", build_mov_or_m4a(handler=b"soun"))
    assert result.supported_for_inspection is True
    assert result.metadata["format"] == "m4a"
    assert result.metadata["duration_seconds"] == pytest.approx(3.0)
    assert result.metadata["has_audio"] is True


@pytest.mark.parametrize(
    ("filename", "bad_bytes"),
    [
        ("junk.wav", b"not a wav"),
        ("junk.flac", b"not a flac stream"),
        ("junk.ogg", b"not an ogg stream"),
    ],
)
async def test_extended_audio_inspection_rejects_garbage(
    storage: LocalStorage, filename, bad_bytes
):
    with pytest.raises(InspectionError):
        await _inspect(storage, filename, bad_bytes)


async def test_audio_m4a_rejects_non_container(storage: LocalStorage):
    with pytest.raises(InspectionError):
        await _inspect(storage, "fake.m4a", b"not a box stream")


# --- extended video ------------------------------------------------------------


async def test_video_mkv_extracts_ebml_metadata(storage: LocalStorage):
    result = await _inspect(storage, "clip.mkv", build_mkv())
    assert result.supported_for_inspection is True
    assert result.media_category == MEDIA_CATEGORY_VIDEO
    assert result.extracted_text is None
    assert result.metadata["format"] == "mkv"
    assert result.metadata["doctype"] == "matroska"
    assert result.metadata["duration_seconds"] == pytest.approx(5.5)
    assert result.metadata["has_video"] is True
    assert result.metadata["has_audio"] is True
    assert result.metadata["width"] == 640
    assert result.metadata["height"] == 480


async def test_video_webm_extracts_doctype_and_codecs(storage: LocalStorage):
    result = await _inspect(storage, "clip.webm", build_mkv(doctype=b"webm"))
    assert result.supported_for_inspection is True
    assert result.metadata["format"] == "webm"
    assert result.metadata["doctype"] == "webm"
    assert result.metadata["codecs"] == "A_OPUS,V_VP9"


async def test_video_avi_extracts_riff_metadata(storage: LocalStorage):
    result = await _inspect(storage, "clip.avi", build_avi())
    assert result.supported_for_inspection is True
    assert result.extracted_text is None
    assert result.metadata["format"] == "avi"
    assert result.metadata["width"] == 320
    assert result.metadata["height"] == 240
    assert result.metadata["fps"] == 25.0
    assert result.metadata["has_video"] is True
    assert result.metadata["has_audio"] is False


async def test_video_mov_extracts_box_metadata(storage: LocalStorage):
    result = await _inspect(storage, "clip.mov", build_mov_or_m4a())
    assert result.supported_for_inspection is True
    assert result.metadata["format"] == "mov"
    assert result.metadata["duration_seconds"] == pytest.approx(3.0)
    assert result.metadata["has_video"] is True


@pytest.mark.parametrize(
    ("filename", "bad_bytes"),
    [
        ("junk.mkv", b"not an EBML stream"),
        ("junk.avi", b"not an AVI RIFF stream"),
        ("junk.webm", b"definitely not webm"),
        ("junk.mov", b"not a quicktime box stream"),
    ],
)
async def test_extended_video_inspection_rejects_garbage(
    storage: LocalStorage, filename, bad_bytes
):
    with pytest.raises(InspectionError):
        await _inspect(storage, filename, bad_bytes)


# --- PPTX ----------------------------------------------------------------------


async def test_pptx_inspection_extracts_slide_text(storage: LocalStorage):
    result = await _inspect(
        storage, "brief.pptx", build_pptx(["Hello slides", "Second deck slide"])
    )

    assert result.supported_for_inspection is True
    assert result.media_category == MEDIA_CATEGORY_PRESENTATION
    assert result.metadata["slide_count"] == 2
    assert result.extracted_text is not None
    assert "Hello slides" in result.extracted_text
    assert "Second deck slide" in result.extracted_text


async def test_pptx_inspection_blank_deck_has_no_text(storage: LocalStorage):
    result = await _inspect(storage, "blank.pptx", build_pptx([]))
    assert result.supported_for_inspection is True
    assert result.metadata["slide_count"] == 0
    assert result.extracted_text is None


async def test_pptx_inspection_malformed_raises_inspection_error(storage: LocalStorage):
    with pytest.raises(InspectionError):
        await _inspect(storage, "broken.pptx", b"this is not a zip")


# --- XLS -----------------------------------------------------------------------


async def test_xls_inspection_extracts_cells(storage: LocalStorage):
    result = await _inspect(storage, "book.xls", build_xls())

    assert result.supported_for_inspection is True
    assert result.media_category == MEDIA_CATEGORY_DOCUMENT
    assert result.metadata["format"] == "xls"
    assert result.metadata["sheet_count"] == 1
    assert result.extracted_text is not None
    assert "host\tcount" in result.extracted_text
    assert "srv-01\t42" in result.extracted_text


async def test_xls_inspection_malformed_raises_inspection_error(storage: LocalStorage):
    with pytest.raises(InspectionError):
        await _inspect(storage, "broken.xls", b"definitely not a BIFF workbook")


# --- EVTX ----------------------------------------------------------------------


async def test_evtx_inspection_extracts_events(storage: LocalStorage):
    data = read_evtx(_EVTX_FIXTURE)
    result = await _inspect(storage, "log.evtx", data)

    assert result.supported_for_inspection is True
    assert result.media_category == MEDIA_CATEGORY_TEXT
    assert result.metadata["format"] == "evtx"
    assert result.metadata["records_processed"] >= 1
    assert result.metadata["records_skipped"] == 0
    assert result.extracted_text is not None
    assert "eventid=4672" in result.extracted_text
    assert "provider=Microsoft-Windows-Security-Auditing" in result.extracted_text


async def test_evtx_inspection_skips_malformed_records_but_stays_supported(storage: LocalStorage):
    data = read_evtx(_EVTX_MALFORMED_FIXTURE)
    result = await _inspect(storage, "dns.evtx", data)

    assert result.supported_for_inspection is True
    assert result.metadata["records_processed"] >= 1
    assert result.metadata["records_skipped"] >= 1


async def test_evtx_inspection_rejects_unreadable_log(storage: LocalStorage):
    with pytest.raises(InspectionError):
        await _inspect(storage, "broken.evtx", b"Hello world, this is not a Windows Event Log")


# --- processing-service integration -------------------------------------------


def _auth(token: str) -> dict[str, str]:
    return {"Authorization": f"Bearer {token}"}


async def _signup_and_login(client: AsyncClient, email: str) -> tuple[str, str]:
    signup = await client.post(
        "/api/v1/auth/signup", json={"email": email, "password": "password123"}
    )
    assert signup.status_code == 201
    login = await client.post(
        "/api/v1/auth/login", json={"email": email, "password": "password123"}
    )
    assert login.status_code == 200
    return login.json()["access_token"]


async def _create_job(client: AsyncClient, token: str) -> str:
    resp = await client.post("/api/v1/jobs", json={"config": {}}, headers=_auth(token))
    assert resp.status_code == 201
    return resp.json()["id"]


async def _upload(
    client: AsyncClient, token: str, job_id: str, filename: str, content: bytes
) -> None:
    content_type = ALLOWED_UPLOAD_TYPES[Path(filename).suffix]
    resp = await client.post(
        f"/api/v1/jobs/{job_id}/input",
        files={"file": (filename, content, content_type)},
        headers=_auth(token),
    )
    assert resp.status_code == 201, resp.text


async def _job_status(job_id: str) -> str:
    async with async_session_factory() as db:
        return await db.scalar(select_job_status(job_id))


def select_job_status(job_id: str):
    from sqlalchemy import select

    return select(Job.status).where(Job.id == uuid.UUID(job_id))


async def _reserve_and_process(job_id: str, storage: LocalStorage):
    async with async_session_factory() as db:
        await reserve_job_for_processing(db, uuid.UUID(job_id))
        return await process_job(uuid.UUID(job_id), db, storage)


async def _read_artifact(storage: LocalStorage, artifact_path: str) -> dict:
    return json.loads(storage.resolve(artifact_path).read_text())


@pytest.mark.parametrize(
    ("filename", "builder"),
    [
        ("plot.svg", lambda: build_svg("<text>alert in svg</text>")),
        ("brief.pptx", lambda: build_pptx(["deck line"])),
        ("book.xls", build_xls),
        ("snd.wav", lambda: build_wav()),
        ("img.webp", build_webp),
        ("clip.mkv", build_mkv),
        ("log.evtx", lambda: read_evtx(_EVTX_FIXTURE)),
    ],
)
async def test_stage_b_format_processes_to_single_artifact(
    client: AsyncClient, storage: LocalStorage, filename, builder
):
    email = unique_email("stgb")
    CREATED_EMAILS.append(email)
    token = await _signup_and_login(client, email)
    job_id = await _create_job(client, token)
    await _upload(client, token, job_id, filename, builder())

    result = await _reserve_and_process(job_id, storage)

    assert await _job_status(job_id) == "completed"
    artifact = await _read_artifact(storage, result.artifact_path)
    entry = artifact["inspected_files"][0]
    assert entry["original_filename"] == filename
    assert entry["supported_for_inspection"] is True
    if filename.endswith((".svg", ".pptx", ".xls", ".evtx")):
        assert entry["extracted_chars"] > 0
    else:
        assert entry["extracted_chars"] is None


@pytest.mark.parametrize(
    ("filename", "bad_bytes"),
    [
        ("broken.svg", b"<svg><text>not xml"),
        ("broken.pptx", b"not a zip"),
        ("broken.xls", b"not BIFF"),
        ("broken.webp", b"junk webp data"),
    ],
)
async def test_stage_b_malformed_format_fails_processing(
    client: AsyncClient, storage: LocalStorage, filename, bad_bytes
):
    email = unique_email("stgb_bad")
    CREATED_EMAILS.append(email)
    token = await _signup_and_login(client, email)
    job_id = await _create_job(client, token)
    await _upload(client, token, job_id, filename, bad_bytes)

    with pytest.raises(ProcessingError):
        await _reserve_and_process(job_id, storage)

    assert await _job_status(job_id) == "failed"