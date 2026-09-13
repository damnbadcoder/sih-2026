from typing import Any
from pptx import Presentation
from app.models.input_file import InputFile
from app.processing.inspection.base import BaseInspector, stream_to_buffer
from app.storage import Storage


class PPTXInspector(BaseInspector):
    media_category = "document"
    supported_extensions = frozenset({".pptx"})
    supported_mime_types = frozenset({
        "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    })

    async def _extract(
        self, file: InputFile, storage: Storage
    ) -> tuple[str | None, dict[str, Any]]:
        buffer = await stream_to_buffer(file, storage)
        try:
            prs = Presentation(buffer)
            text_lines = []
            slide_count = len(prs.slides)

            for idx, slide in enumerate(prs.slides, start=1):
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for paragraph in shape.text_frame.paragraphs:
                            text = paragraph.text.strip()
                            if text:
                                text_lines.append(f"[Slide {idx}] {text}")

                if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                    notes = slide.notes_slide.notes_text_frame.text.strip()
                    if notes:
                        text_lines.append(f"[Slide {idx} Notes] {notes}")

            extracted_text = "\n".join(text_lines) if text_lines else None
            metadata = {
                "slide_count": slide_count,
                "format": "pptx",
            }
            return extracted_text, metadata
        finally:
            buffer.close()